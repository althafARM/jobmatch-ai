import os
import re
from io import BytesIO

from docx import Document
from pypdf import PdfReader
from PIL import Image, ImageOps
import pytesseract


# --------------------------------------------------
# 1. Configure Tesseract OCR
# --------------------------------------------------

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

if os.path.isfile(TESSERACT_PATH):
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# --------------------------------------------------
# 2. Skills supported by this version
# --------------------------------------------------

KNOWN_SKILLS = [
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "data analysis",
    "data science",
    "scikit-learn",
    "spring boot",
    "javascript",
    "typescript",
    "node.js",
    "fastapi",
    "django",
    "flask",
    "python",
    "java",
    "react",
    "sql",
    "mysql",
    "postgresql",
    "mongodb",
    "pandas",
    "numpy",
    "tensorflow",
    "pytorch",
    "excel",
    "power bi",
    "tableau",
    "docker",
    "git",
    "html",
    "css",
    "aws",
    "c++",
    "c#",
]


# --------------------------------------------------
# 3. Extract text from an image using OCR
# --------------------------------------------------

def extract_image_text(content: bytes) -> str:
    """Extract text from JPG, JPEG, or PNG images."""

    if not os.path.isfile(TESSERACT_PATH):
        raise RuntimeError(
            "Tesseract OCR was not found at: "
            f"{TESSERACT_PATH}"
        )

    try:
        with Image.open(BytesIO(content)) as original_image:
            image = ImageOps.exif_transpose(original_image)
            image = image.convert("RGB")

            # Enlarge small images to improve OCR.
            if image.width < 1500:
                ratio = 1500 / image.width
                new_size = (
                    1500,
                    int(image.height * ratio),
                )
                image = image.resize(
                    new_size,
                    Image.Resampling.LANCZOS,
                )

            extracted_text = pytesseract.image_to_string(
                image,
                lang="eng",
                config="--psm 6",
            )

        if not extracted_text.strip():
            raise ValueError(
                "No readable text was found in the image. "
                "Try a clearer image of your resume."
            )

        return extracted_text.strip()

    except pytesseract.TesseractNotFoundError as error:
        raise RuntimeError(
            "Tesseract OCR could not be started. "
            "Check the Tesseract installation path."
        ) from error

    except pytesseract.TesseractError as error:
        raise RuntimeError(
            f"Tesseract OCR failed: {error}"
        ) from error

    except (OSError, ValueError) as error:
        raise ValueError(
            f"Could not read the resume image: {error}"
        ) from error


# --------------------------------------------------
# 4. Extract text from PDF, DOCX, and image resumes
# --------------------------------------------------

def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from supported resume formats."""

    name = os.path.basename(filename).lower().strip()

    if not content:
        raise ValueError("The uploaded resume is empty.")

    # Image resumes
    if name.endswith((".png", ".jpg", ".jpeg")):
        return extract_image_text(content)

    # PDF resumes
    if name.endswith(".pdf"):
        try:
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise ValueError(
                    "Encrypted PDFs are not supported."
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            ).strip()

            # Scanned PDFs may need a separate PDF OCR step.
            if not text:
                raise ValueError(
                    "No selectable text was found in this PDF. "
                    "It may be a scanned PDF."
                )

            return text

        except ValueError:
            raise

        except Exception as error:
            raise ValueError(
                "Could not read the PDF resume."
            ) from error

    # DOCX resumes
    if name.endswith(".docx"):
        try:
            document = Document(BytesIO(content))
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            # Extract text from tables too.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.extend(
                        cell.text for cell in row.cells
                    )

            text = "\n".join(paragraphs).strip()

            if not text:
                raise ValueError(
                    "No readable text was found in the DOCX file."
                )

            return text

        except ValueError:
            raise

        except Exception as error:
            raise ValueError(
                "Could not read the DOCX resume."
            ) from error

    raise ValueError(
        "Unsupported file format. "
        "Upload a PDF, DOCX, JPG, JPEG, or PNG resume."
    )


# --------------------------------------------------
# 5. Extract known skills from resume text
# --------------------------------------------------

def extract_skills(text: str) -> list[str]:
    """Find supported skills mentioned in resume text."""

    normalized_text = text.lower()
    found_skills = []

    for skill in KNOWN_SKILLS:
        pattern = (
            r"(?<!\w)"
            + re.escape(skill.lower())
            + r"(?!\w)"
        )

        if re.search(pattern, normalized_text):
            found_skills.append(skill)

    return found_skills